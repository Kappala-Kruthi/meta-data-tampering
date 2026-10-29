"""
reports/pptx_exporter.py
Generates an editable 14-slide academic PowerPoint presentation (.pptx).
Designed for college project demonstration and viva evaluation.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def create_academic_presentation(output_path: str) -> str:
    """
    Builds a professional 14-slide PowerPoint presentation for college viva demonstration.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_BG = RGBColor(0x0A, 0x0E, 0x1A)      # Deep Dark Navy
    COLOR_CARD = RGBColor(0x16, 0x1F, 0x30)    # Slate Card
    COLOR_TEXT = RGBColor(0xF8, 0xFA, 0xFC)    # Crisp White
    COLOR_MUTED = RGBColor(0x94, 0xA3, 0xB8)   # Light Slate
    COLOR_CYAN = RGBColor(0x06, 0xB6, 0xD4)    # Forensic Cyan
    COLOR_INDIGO = RGBColor(0x63, 0x66, 0xF1)  # Accent Indigo

    blank_layout = prs.slide_layouts[6]  # blank layout

    def add_styled_slide(title_text: str, subtitle_text: str = ""):
        slide = prs.slides.add_slide(blank_layout)
        
        # Dark Background shape
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))  # 1 = MSO_SHAPE.RECTANGLE
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        # Title Box
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_CYAN

        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = COLOR_MUTED
            p_sub.space_before = Pt(4)

        return slide

    def fill_bullet_card(slide, bullets, top_offset=1.8, card_height=5.0):
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(top_offset), Inches(11.7), Inches(card_height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_INDIGO
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.35)

        for i, (title, desc) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            
            r1 = p.add_run()
            r1.text = f"• {title}: "
            r1.font.bold = True
            r1.font.size = Pt(14)
            r1.font.color.rgb = COLOR_CYAN

            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(13)
            r2.font.color.rgb = COLOR_TEXT

    # ---------------- SLIDE 1: TITLE SLIDE ----------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "ADVANCED DIGITAL FORENSICS SYSTEM"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p2 = tf1.add_paragraph()
    p2.text = "Metadata Tampering Detection in Media Files Using ExifTool"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT
    p2.space_before = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Candidate: Kappala Bhuvana Kruthi (bhuvanakruthi.cutm@gmail.com)\nDepartment of Computer Science & Engineering | Academic Year 2026"
    p3.font.size = Pt(16)
    p3.font.color.rgb = COLOR_MUTED
    p3.space_before = Pt(24)

    # ---------------- SLIDE 2: PROBLEM STATEMENT ----------------
    s2 = add_styled_slide("1. Problem Statement", "Susceptibility of digital media evidence to invisible metadata alterations")
    fill_bullet_card(s2, [
        ("Digital Evidence Vulnerability", "Digital photos and media files are widely used in legal proceedings, journalism, and insurance claims. However, metadata (EXIF/XMP) can be silently modified without leaving visible image artifacts."),
        ("Limitations of Built-in OS Tools", "Operating systems display less than 10% of available metadata tags and cannot parse complex MakerNotes, GPS history, or XMP modification records."),
        ("Privacy & Chain of Custody Violations", "Existing online tools force investigators to upload evidence to third-party cloud servers, violating privacy laws and invalidating legal chain-of-custody."),
        ("Need for Automated Differential Analysis", "Forensic examiners need an offline, explainable system to rapidly detect alterations between reference and suspected images."),
    ])

    # ---------------- SLIDE 3: PROJECT OBJECTIVES ----------------
    s3 = add_styled_slide("2. Project Objectives", "Key capabilities designed into the forensics system")
    fill_bullet_card(s3, [
        ("Subprocess ExifTool Integration", "Safely execute Phil Harvey's ExifTool standalone binary via Python subprocess (shell=False) to extract comprehensive metadata across all tag families."),
        ("Cryptographic Integrity Verification", "Compute bit-for-bit SHA-256 and legacy MD5 checksums to verify mathematical byte integrity."),
        ("Two-File Differential Analysis", "Automate comparison between File A (Reference) and File B (Suspected) to classify tags as Matching, Changed, Only A, or Only B."),
        ("Explainable Anomaly Rules", "Implement transparent heuristics to detect editing software signatures (Photoshop/GIMP), chronological timestamp inversions, and GPS stripping."),
        ("Offline Local Persistence & Reporting", "Store audit trails in a local SQLite database and export reports in HTML, JSON, CSV, and DOCX formats."),
    ])

    # ---------------- SLIDE 4: DIGITAL FORENSICS & METADATA BASICS ----------------
    s4 = add_styled_slide("3. Digital Forensics & Metadata Fundamentals", "Understanding hidden layers within media files")
    fill_bullet_card(s4, [
        ("EXIF (Exchangeable Image File Format)", "Stores camera hardware settings: Make, Model, Shutter Speed, Aperture, ISO, Lens Serial Number, and Original Capture Date."),
        ("XMP (Extensible Metadata Platform)", "Adobe-developed XML metadata standard. Tracks modification history, editing software agents, document IDs, and derivate records."),
        ("IPTC-NAA Standard", "Captures journalistic and archival information: Copyright, Author, Location Details, Captions, and Keywords."),
        ("GPS Geolocation Group", "Records physical recording location: Latitude, Longitude, Altitude, Timestamp, and Map Positioning."),
        ("MakerNotes", "Proprietary camera manufacturer binary blocks (Canon, Nikon, Sony) often stripped or corrupted during non-camera editing."),
    ])

    # ---------------- SLIDE 5: WHY EXIFTOOL ----------------
    s5 = add_styled_slide("4. ExifTool: The Forensic Standard", "Why Phil Harvey's ExifTool is superior to basic libraries")
    fill_bullet_card(s5, [
        ("Gold Standard Tooling", "Phil Harvey's ExifTool is recognized globally by law enforcement and cyber forensic laboratories as the premier metadata reader and writer."),
        ("Comprehensive Tag Support", "Reads thousands of tags across 100+ file formats including JPEG, PNG, TIFF, WebP, HEIC, and RAW camera formats."),
        ("Safe JSON Extraction", "Using `-j -G -a -u`, ExifTool exports structured JSON preserving tag families and unknown vendor tags without corrupting source files."),
        ("No File Locking / Read-Only Mode", "ExifTool operates in read-only forensic mode, ensuring that original evidence files are never modified during analysis."),
    ])

    # ---------------- SLIDE 6: SYSTEM ARCHITECTURE ----------------
    s6 = add_styled_slide("5. System Architecture", "3-Tier modular design ensuring forensic integrity and security")
    fill_bullet_card(s6, [
        ("Presentation Tier", "Modern Cyber-Forensics Dark UI built with Flask templates, Vanilla CSS, and Vanilla JavaScript. Zero external trackers or CDNs."),
        ("Core Forensic Engine", "Modular Python engine:\n• exiftool_runner: Subprocess controller\n• hasher: Chunked SHA-256/MD5 streamer\n• comparator: Differential comparison engine\n• anomaly_detector: Rule-based heuristics evaluator"),
        ("Persistence & Native Tools", "Standalone ExifTool 64-bit binary running locally in project tools/ folder; SQLite database (investigations.db) storing case records and chain-of-custody audit logs."),
    ])

    # ---------------- SLIDE 7: INVESTIGATION WORKFLOW ----------------
    s7 = add_styled_slide("6. Forensic Analysis Workflow", "Step-by-step pipeline from evidence ingestion to formal report")
    fill_bullet_card(s7, [
        ("Evidence Ingestion & Validation", "Uploaded files are assigned randomized server UUIDs and checked against MIME allowlists (JPEG, PNG, WebP, TIFF, HEIC)."),
        ("Cryptographic Hashing", "SHA-256 and MD5 hashes are calculated in 64 KB memory chunks before any metadata extraction begins."),
        ("ExifTool Execution", "Subprocess executes standalone ExifTool with `-j -G -a -u`, extracting all tag families and formatting timestamps."),
        ("Heuristic Anomaly Evaluation", "Scanners inspect for software signatures, timestamp contradictions, GPS presence/stripping, and format mismatches."),
        ("Differential Categorization & Audit Logging", "If comparing two files, differential engine groups tags; records are stored in SQLite and reports generated."),
    ])

    # ---------------- SLIDE 8: ANOMALY HEURISTICS ----------------
    s8 = add_styled_slide("7. Tampering Heuristic Rules", "Transparent, rule-based detection without fabricated AI percentages")
    fill_bullet_card(s8, [
        ("Editing Software Signature (ANOM_EDITOR_SIGNATURE)", "Detects traces of editing suites (Adobe Photoshop, GIMP, Canva, Lightroom, Snapseed) in Software, CreatorTool, or History tags."),
        ("Chronological Timestamp Inversion (ANOM_TIMESTAMP_INVERSION)", "Flags cases where ModifyDate is earlier than DateTimeOriginal (physically impossible under standard camera operations)."),
        ("Extension / Format Discrepancy (ANOM_EXT_TYPE_MISMATCH)", "Flags instances where the file extension (e.g., .jpg) conflicts with actual internal binary structure (e.g., PNG)."),
        ("Metadata Stripping / Sparsity (ANOM_METADATA_STRIPPED)", "Notes files lacking camera hardware details or having fewer than 15 tags (common in social media downloads or privacy scrubbing)."),
        ("Differential Injections & Stripping", "Detects when File B introduces editor software tags or removes GPS coordinates present in File A."),
    ])

    # ---------------- SLIDE 9: HASH INTEGRITY ----------------
    s9 = add_styled_slide("8. Cryptographic Hash Verification", "Mathematical certainty of bit-for-bit file integrity")
    fill_bullet_card(s9, [
        ("SHA-256 (Primary Forensic Standard)", "256-bit secure cryptographic hash. Changing even a single bit in an image completely alters the resulting 64-hexadecimal character digest (Avalanche Effect)."),
        ("MD5 (Legacy Compatibility)", "Included solely for cross-referencing legacy police and forensic databases. Explicitly marked in the UI as cryptographically broken for collision resistance."),
        ("Identical Hashes = Exact Byte Duplicate", "If two files share the same SHA-256 hash, they are identical in every single byte with astronomical mathematical confidence."),
        ("Crucial Forensic Principle", "A differing hash proves binary non-identity, but DOES NOT alone establish malicious manipulation (re-compression or rotation also changes hashes)."),
    ])

    # ---------------- SLIDE 10: TWO-FILE WORKSPACE ----------------
    s10 = add_styled_slide("9. Two-File Forensic Comparison Workspace", "Structured differential matrix between Reference (File A) and Suspected (File B)")
    fill_bullet_card(s10, [
        ("Side-by-Side Metadata Alignment", "Extracts metadata from both media files and aligns them key-by-key across all tag families."),
        ("Four Differential Categories", "• MATCHING: Same tag, normalized equal values\n• CHANGED: Tag present in both, but values altered\n• ONLY IN FILE A: Tag stripped or absent in File B\n• ONLY IN FILE B: Tag newly introduced in File B"),
        ("Similarity Percentage Calculation", "Computes objective tag similarity metric based on matching vs total tags."),
        ("Instant Filtering & Search", "Examiners can filter by category tab or use the live search bar to inspect specific tags like 'Software', 'GPS', or 'DateTime'."),
        ("Multi-Format Export", "Supports one-click export of comparison tables to CSV, JSON, and printable HTML."),
    ])

    # ---------------- SLIDE 11: DEMO & TEST RESULTS ----------------
    s11 = add_styled_slide("10. Automated Validation & Test Suite", "Rigorous verification: 8 of 8 tests executed and passed")
    fill_bullet_card(s11, [
        ("TC-01: Cryptographic Hasher (PASSED)", "Validated bit-for-bit match on duplicates and mismatch upon altering 1 byte."),
        ("TC-02: ExifTool Binary Detection (PASSED)", "Verified standalone ExifTool 13.59 executable launches and returns version."),
        ("TC-03: Comparator Differential Logic (PASSED)", "Verified classification into Matching, Changed, Only A, and Only B."),
        ("TC-04: Anomaly Heuristics Scanner (PASSED)", "Verified detection of Adobe Photoshop signature and timestamp inversion."),
        ("TC-05: SQLite Persistence & Audit Log (PASSED)", "Verified case saving, retrieval, listing, and audit event recording."),
        ("TC-06 to TC-08: Real ExifTool Extraction on Sample Images (PASSED)", "Verified end-to-end extraction from genuine synthetic camera, edited, and stripped images."),
    ])

    # ---------------- SLIDE 12: FORENSIC LIMITATIONS ----------------
    s12 = add_styled_slide("11. Forensic Principles & Limitations", "Maintaining scientific objectivity in digital evidence analysis")
    fill_bullet_card(s12, [
        ("Hashes Detect Changes, Not Motives", "A hash difference proves non-identity; it cannot prove malicious forgery versus benign format re-encoding."),
        ("Metadata Stripping is Widespread", "Social media networks (WhatsApp, Instagram, X) strip metadata automatically for privacy; lack of metadata does not equal guilt."),
        ("Device Clocks are Unverified", "EXIF timestamps reflect the camera's internal clock at recording time, which may be uncalibrated or set to an incorrect timezone."),
        ("Corroboration is Essential", "Digital forensics requires correlating metadata anomalies with visual inspection, file system timestamps, and external testimony."),
    ])

    # ---------------- SLIDE 13: FUTURE SCOPE ----------------
    s13 = add_styled_slide("12. Future Scope & Enhancements", "Expanding capabilities for advanced forensic workflows")
    fill_bullet_card(s13, [
        ("Pixel-Level Error Level Analysis (ELA)", "Integrate compression artifact analysis to visually detect resaved composite areas within pixel grids."),
        ("Multimedia Video/Audio Inspection", "Extend differential extraction to MP4, MKV, and WAV container structures."),
        ("PRNU Sensor Noise Fingerprinting", "Identify the unique physical camera sensor that captured an image via Photo-Response Non-Uniformity."),
        ("Integration with Forensic Evidence Standards", "Support automated packaging into standard DFIR forensic container formats (e.g., AFF4 / E01)."),
    ])

    # ---------------- SLIDE 14: CONCLUSION ----------------
    s14 = add_styled_slide("13. Project Conclusion & Summary", "Summary of project achievements")
    fill_bullet_card(s14, [
        ("Applied Academic Achievement", "Successfully engineered a fully working, offline digital forensics application capable of real ExifTool execution, differential comparison, and cryptographic verification."),
        ("Defensible & Objective Methodology", "Implements transparent heuristic rules and rigorous caveat disclosures adhering to ISO/IEC 27037 standards."),
        ("Comprehensive Deliverables", "Includes web dashboard, automated test suite, sample test images, formal academic report, and presentation slides."),
        ("Open for Questions", "Thank you! I am ready for questions from the evaluation committee."),
    ])

    prs.save(output_path)
    return output_path
